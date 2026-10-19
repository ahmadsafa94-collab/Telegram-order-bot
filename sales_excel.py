"""Builds the daily sales workbook the admin gets from the bot.

Kept apart from bot.py so it only deals with plain data: bot.py collects
the paid orders and the per-product costs, this file lays them out.

Sheets:
  Daily          one row per day: orders, accounts, received, costs, profit
  Day details    every product sold each day with its cost
  Accounts sold  accounts per product with their share, plus a pie chart
  Costs          the cost per account currently set in the bot
"""

import io

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties

FONT = "Arial"
USD = '$#,##0.00;($#,##0.00);"-"'
DATE = "ddd dd mmm yyyy"
PCT = '0.0%;-0.0%;"-"'
HEADER_FILL = PatternFill("solid", start_color="0D9488")
TOTAL_FILL = PatternFill("solid", start_color="E6F4F2")
THIN = Side(style="thin", color="B7C9C6")
# Inputs the owner might want to check (costs, received) are blue, formulas black.
INPUT_FONT = Font(name=FONT, color="0000FF")
BODY_FONT = Font(name=FONT)
BOLD_FONT = Font(name=FONT, bold=True)
# Pie slices beyond this many products are grouped into "Other".
PIE_SLICES = 8


def _title(ws, text, subtitle):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=FONT, italic=True, color="666666")


def _header(ws, row, labels, widths):
    for col, (label, width) in enumerate(zip(labels, widths), start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _style_row(ws, row, ncols, fmts, fonts=None, fill=None):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = (fonts or {}).get(col, BODY_FONT)
        cell.border = Border(bottom=THIN)
        if col in fmts:
            cell.number_format = fmts[col]
        if fill:
            cell.fill = fill


def build_sales_workbook(days, products_sold, costs, generated_label: str) -> bytes:
    """days: [{"day": date, "orders": int, "received": float,
               "items": [(product name, accounts, cost per account)]}], oldest first.
    products_sold: [(product name, accounts)] over all days, most sold first.
    costs: [(product name, item id, cost per account)].
    Returns the .xlsx file as bytes."""
    wb = Workbook()

    # ---- Day details (built first: Daily sums from it) ----
    det = wb.active
    det.title = "Day details"
    _title(det, "Accounts sold per day", "Cost per account is what was set in the bot when the order was paid.")
    _header(det, 4, ["Date", "Product", "Accounts", "Cost per account (USD)", "Cost (USD)"],
            [16, 58, 11, 14, 13])
    r = 5
    for d in days:
        for name, qty, cost_each in d["items"]:
            det.cell(row=r, column=1, value=d["day"])
            det.cell(row=r, column=2, value=name)
            det.cell(row=r, column=3, value=qty)
            det.cell(row=r, column=4, value=cost_each)
            det.cell(row=r, column=5, value=f"=C{r}*D{r}")
            _style_row(det, r, 5, {1: DATE, 4: USD, 5: USD}, fonts={4: INPUT_FONT})
            r += 1
    det_last = max(r - 1, 5)

    # ---- Daily ----
    ws = wb.create_sheet("Daily", 0)
    _title(ws, "Medic Accounts — daily sales (USD)", f"Updated {generated_label}. A new row is added every day.")
    _header(ws, 4, ["Date", "Orders", "Accounts sold", "Received (USD)", "Costs (USD)", "Profit (USD)"],
            [16, 9, 11, 14, 13, 13])
    ws["D4"].comment = Comment("Amount customers actually paid, after discounts and credits.", "Bot")
    ws["E4"].comment = Comment("Accounts sold × cost per account (from the Day details sheet).", "Bot")
    first = 5
    r = first
    for d in days:
        ws.cell(row=r, column=1, value=d["day"])
        ws.cell(row=r, column=2, value=d["orders"])
        ws.cell(row=r, column=3, value=f"=SUMIF('Day details'!$A$5:$A${det_last},A{r},'Day details'!$C$5:$C${det_last})")
        ws.cell(row=r, column=4, value=round(d["received"], 2))
        ws.cell(row=r, column=5, value=f"=SUMIF('Day details'!$A$5:$A${det_last},A{r},'Day details'!$E$5:$E${det_last})")
        ws.cell(row=r, column=6, value=f"=D{r}-E{r}")
        _style_row(ws, r, 6, {1: DATE, 4: USD, 5: USD, 6: USD},
                   fonts={2: INPUT_FONT, 4: INPUT_FONT, 6: BOLD_FONT})
        r += 1
    last = max(r - 1, first)
    ws.cell(row=r, column=1, value="Total")
    for col, letter in ((2, "B"), (3, "C"), (4, "D"), (5, "E"), (6, "F")):
        ws.cell(row=r, column=col, value=f"=SUM({letter}{first}:{letter}{last})")
    _style_row(ws, r, 6, {4: USD, 5: USD, 6: USD}, fonts={c: BOLD_FONT for c in range(1, 7)}, fill=TOTAL_FILL)

    # ---- Accounts sold + pie chart ----
    acc = wb.create_sheet("Accounts sold", 1)
    _title(acc, "Accounts sold per product", "All days in this file.")
    _header(acc, 4, ["Product", "Accounts", "Share"], [58, 11, 10])
    r = 5
    for name, qty in products_sold:
        acc.cell(row=r, column=1, value=name)
        acc.cell(row=r, column=2, value=qty)
        r += 1
    acc_last = max(r - 1, 5)
    total_row = r
    for rr in range(5, total_row):
        acc.cell(row=rr, column=3, value=f"=IF($B${total_row}=0,0,B{rr}/$B${total_row})")
        _style_row(acc, rr, 3, {3: PCT})
    acc.cell(row=total_row, column=1, value="Total")
    acc.cell(row=total_row, column=2, value=f"=SUM(B5:B{acc_last})")
    acc.cell(row=total_row, column=3, value=f"=SUM(C5:C{acc_last})")
    _style_row(acc, total_row, 3, {3: PCT}, fonts={c: BOLD_FONT for c in range(1, 4)}, fill=TOTAL_FILL)

    if products_sold:
        # The chart reads a small block of its own so the pie stays readable:
        # the biggest products, then everything else as "Other".
        top = products_sold[:PIE_SLICES]
        rest = sum(q for _, q in products_sold[PIE_SLICES:])
        slices = top + ([("Other", rest)] if rest else [])
        acc["E4"] = "Chart data"
        acc["E4"].font = Font(name=FONT, bold=True, color="666666")
        acc.column_dimensions["E"].width = 34
        acc.column_dimensions["F"].width = 9
        for i, (name, qty) in enumerate(slices):
            short = name if len(name) <= 34 else name[:33] + "…"
            acc.cell(row=5 + i, column=5, value=short).font = Font(name=FONT, color="666666")
            acc.cell(row=5 + i, column=6, value=qty).font = Font(name=FONT, color="666666")
        end = 5 + len(slices) - 1

        pie = PieChart()
        pie.title = "Accounts sold"
        pie.add_data(Reference(acc, min_col=6, min_row=5, max_row=end))
        pie.set_categories(Reference(acc, min_col=5, min_row=5, max_row=end))
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showVal = True
        pie.dataLabels.showPercent = True
        pie.dataLabels.showCatName = False
        pie.dataLabels.showSerName = False
        pie.dataLabels.showLeaderLines = True
        pie.height = 11
        pie.width = 18
        acc.add_chart(pie, f"A{total_row + 3}")

    # ---- Costs ----
    cs = wb.create_sheet("Costs")
    _title(cs, "Cost per account (USD)",
           "To change a cost: in the bot, 📦 Manage Stock → tap the product → 💲 Edit Cost.")
    _header(cs, 4, ["Product", "Item ID", "Cost per account (USD)"], [58, 24, 14])
    for i, (name, item_id, cost) in enumerate(costs, start=5):
        cs.cell(row=i, column=1, value=name)
        cs.cell(row=i, column=2, value=item_id)
        cs.cell(row=i, column=3, value=cost)
        _style_row(cs, i, 3, {3: USD}, fonts={3: INPUT_FONT})

    # The bot can't compute formulas itself; ask the app to do it on open.
    wb.calculation = CalcProperties(fullCalcOnLoad=True)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
